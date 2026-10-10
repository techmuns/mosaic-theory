"""Write the audit's deliverables: CSVs, the coverage and validation reports, and the updated
workbook (a COPY of the input — the input file is never opened for writing).

WORKBOOK CONVENTIONS are the input's own: Aptos, navy header bands, blue-on-cream ``[INPUT]``
cells, green cross-sheet links. Filled Forward-Return-Test inputs LINK to their source rows
(``'Polygon Daily Prices'!E…`` for the stock's adjClose AND for URTH, the USD world benchmark
PROXY) so the sheet's existing R–U formulas recompute from traceable cells, as the pilot rows do.
The pilot rows themselves are left exactly as supplied; their currency finding is stated beside
them, never written over them.
"""

from __future__ import annotations

import csv
import dataclasses
import json
import shutil
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from vinci.research_audit.engine import COMPLETE, Identity, Obs
from vinci.research_audit.prices import Series
from vinci.research_audit.workbook import SHEET_TESTS, AuditWorkbook

FONT = "Aptos"
NAVY, TITLE_FILL, INPUT_FILL = "213D54", "13253A", "FFF8DE"
PCT = "0.0%;[Red]-0.0%;0.0%"
LVL = "0.00"


# --- CSV ------------------------------------------------------------------------------------------

def _cell(v: Any) -> Any:
    if isinstance(v, date | datetime):
        return v.isoformat()
    if isinstance(v, list | tuple):
        return " | ".join(str(x) for x in v)
    return "" if v is None else v


def columns(rows: Sequence[dict[str, Any]]) -> list[str]:
    """Ordered UNION of every row's keys — a mixed table never loses a column to its first row."""
    seen: dict[str, None] = {}
    for r in rows:
        for k in r:
            seen.setdefault(k, None)
    return list(seen)


def write_csv(path: Path, rows: Sequence[dict[str, Any]], headers: Sequence[str] | None = None
              ) -> None:
    hdr = list(headers or columns(rows))
    with path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=hdr, extrasaction="raise")
        w.writeheader()
        for r in rows:
            w.writerow({k: _cell(r.get(k)) for k in hdr})


def obs_rows(ledger: Iterable[Obs]) -> list[dict[str, Any]]:
    return [dataclasses.asdict(o) for o in ledger]


def price_rows(series: dict[str, Series], audit_of: dict[str, str]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for sym in sorted(series):
        for p in series[sym].points:
            out.append({
                "audit_ticker": audit_of.get(sym, sym), "provider_ticker": sym,
                "date": p.day.isoformat(), "open": p.open, "high": p.high, "low": p.low,
                "close": p.close, "adjClose": p.adj_close, "divCash": p.div_cash,
                "splitFactor": p.split_factor, "volume": p.volume, "source": p.source,
                "source_ref": p.source_ref,
            })
    return out


# --- workbook helpers -----------------------------------------------------------------------------

def _hdr(ws: Worksheet, row: int, headers: Sequence[str], widths: dict[int, float] | None = None
         ) -> None:
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name=f"{FONT} Display", size=16, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TITLE_FILL)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = Font(name=FONT, size=10, italic=True)
    s.alignment = Alignment(wrap_text=False)


def _put(ws: Worksheet, row: int, col: int, v: Any, fmt: str | None = None,
         color: str | None = None, fill: str | None = None, bold: bool = False) -> None:
    c = ws.cell(row=row, column=col, value=_xl(v))
    c.font = Font(name=FONT, size=10, color=color, bold=bold)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)


def _xl(v: Any) -> Any:
    if isinstance(v, list | tuple):
        return " | ".join(str(x) for x in v)
    if isinstance(v, str) and v[:4].isdigit() and len(v) == 10 and v[4] == "-" and v[7] == "-":
        return date.fromisoformat(v)
    return v


def _fmt_for(key: str) -> str | None:
    k = key.lower()
    if k.endswith("_pp") or "_pp_" in k:
        return '0.00" pp"'  # already in percentage points — never re-scaled as a fraction
    if "win_rate" in k or "share_" in k:
        return "0%"
    if any(t in k for t in ("return", "excess", "median", "mean", "p10", "p25", "p75", "p90",
                            "min", "max", "stdev", "_tr", "diff_", "_1d")):
        return PCT
    if k.endswith(("_entry", "_exit", "adj_close", "close")):
        return LVL
    return None


def table_sheet(xl: openpyxl.Workbook, name: str, title: str, subtitle: str,
                rows: Sequence[dict[str, Any]], headers: Sequence[str] | None = None,
                start: int = 4) -> Worksheet:
    ws: Worksheet = xl.create_sheet(name)
    _title(ws, title, subtitle)
    hdr = list(headers or columns(rows) or ["(no rows)"])
    _hdr(ws, start, hdr)
    for i, r in enumerate(rows, start + 1):
        for j, k in enumerate(hdr, 1):
            _put(ws, i, j, r.get(k), _fmt_for(k))
    ws.freeze_panes = f"B{start + 1}"
    return ws


# --- the updated workbook -------------------------------------------------------------------------

def write_workbook(src: Path, dst: Path, wb: AuditWorkbook, ledger: list[Obs],
                   series: dict[str, Series], audit_of: dict[str, str],
                   sheets: dict[str, tuple[str, list[dict[str, Any]]]],
                   summary_lines: list[tuple[str, str]], world_symbol: str) -> dict[str, int]:
    """Copy ``src`` → ``dst`` and update the COPY. Returns counts of what was filled."""
    shutil.copyfile(src, dst)
    xl = openpyxl.load_workbook(dst)  # formulas kept (no data_only)

    # 1. Polygon Daily Prices — the link target for every filled stock level.
    ps = xl.create_sheet("Polygon Daily Prices")
    _title(ps, "POLYGON DAILY PRICES | RAW CLOSE → TOTAL-RETURN adjClose",
           "close = as-traded; adjClose = dividends reinvested at ex-date close, split-adjusted, "
           "scaled to equal close on the last bar; divCash / splitFactor per session. Sources per "
           "row (live aggregates adjusted=false, or the repo's grouped-daily cache un-adjusted).")
    cols = ["audit_ticker", "provider_ticker", "date", "close", "adjClose", "divCash",
            "splitFactor", "open", "high", "low", "volume", "source"]
    _hdr(ps, 4, cols, {12: 48})
    where: dict[tuple[str, date], int] = {}
    r = 5
    for row in price_rows(series, audit_of):
        for j, k in enumerate(cols, 1):
            v = row[k]
            c = ps.cell(row=r, column=j, value=date.fromisoformat(v) if k == "date" else v)
            if k == "date":
                c.number_format = "yyyy-mm-dd"
        where[(row["provider_ticker"], date.fromisoformat(row["date"]))] = r
        r += 1
    ps.freeze_panes = "C5"

    # 2. Fill Forward Return Tests inputs (pilot rows untouched).
    ft = xl[SHEET_TESTS]
    by_id = {o.test_id: o for o in ledger if o.test_id}
    extra = ["Audit status [PIPELINE]", "Scoring class", "Excess vs S&P 500 Net TR (CIQ)",
             "Sector benchmark (PROXY)", "Excess vs sector PROXY", "Polygon ticker", "Flags",
             "USD excess vs URTH PROXY [PIPELINE]"]
    for j, h in enumerate(extra, 25):
        c = ft.cell(row=4, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True)
        ft.column_dimensions[get_column_letter(j)].width = 50 if j == 31 else 22
    filled = untouched_pilot = not_computed = 0
    blue, green = "0000FF", "008000"
    for t in wb.tests:
        o = by_id.get(t.test_id)
        scoring = {"LONG": "LONG — scored", "NEUTRAL": "EQUAL-WEIGHT — reported separately",
                   "UNDERWEIGHT": "UNDERWEIGHT — reported separately, not scored as a short"}
        if t.populated:
            untouched_pilot += 1
            _put(ft, t.row, 25, "CIQ PILOT (unchanged). Stock TR reproduced from Polygon; its "
                 "benchmark column is NOT USD (CIQ series fits EUR) — USD excess in column AF; "
                 "see 'Validation'")
            if o:
                _put(ft, t.row, 26, scoring.get(o.direction, o.direction))
                _put(ft, t.row, 32, o.excess_vs_world, PCT)
            continue
        if o is None:
            _put(ft, t.row, 25, "NOT IN LEDGER — check")
            continue
        _put(ft, t.row, 26, scoring.get(o.direction, o.direction))
        _put(ft, t.row, 30, o.provider_ticker)
        _put(ft, t.row, 31, "; ".join(o.flags))
        if o.status != COMPLETE:
            not_computed += 1
            _put(ft, t.row, 25, f"{o.status}: {o.status_detail}")
            _put(ft, t.row, 23, f"{o.status} — no inputs filled (see Coverage & Unresolved)",
                 color=blue, fill=INPUT_FILL)
            continue
        e, x = date.fromisoformat(o.entry_date), date.fromisoformat(o.exit_date)
        _put(ft, t.row, 12, e, "dd-mmm-yyyy", blue, INPUT_FILL)
        _put(ft, t.row, 13, x, "dd-mmm-yyyy", blue, INPUT_FILL)
        pe, px = where[(o.provider_ticker, e)], where[(o.provider_ticker, x)]
        be, bx = where[(world_symbol, e)], where[(world_symbol, x)]
        _put(ft, t.row, 14, f"='Polygon Daily Prices'!E{pe}", LVL, green, INPUT_FILL)
        _put(ft, t.row, 15, f"='Polygon Daily Prices'!E{px}", LVL, green, INPUT_FILL)
        _put(ft, t.row, 16, f"='Polygon Daily Prices'!E{be}", LVL, green, INPUT_FILL)
        _put(ft, t.row, 17, f"='Polygon Daily Prices'!E{bx}", LVL, green, INPUT_FILL)
        _put(ft, t.row, 22, f"{world_symbol} total return — PROXY for MSCI World Net USD (the "
                            f"CIQ series is not USD; see 'Benchmarks & Tracking')")
        _put(ft, t.row, 23,
             f"Stock: Polygon {o.provider_ticker} adjClose ('Polygon Daily Prices' rows {pe}/{px}"
             f"; raw close + Polygon dividends/splits). Benchmark: Polygon {world_symbol} "
             f"adjClose (rows {be}/{bx}). Pipeline: vinci.research_audit.",
             color=blue, fill=INPUT_FILL)
        _put(ft, t.row, 32, o.excess_vs_world, PCT)
        _put(ft, t.row, 25, "COMPLETE")
        _put(ft, t.row, 27, o.excess_vs_sp500, PCT)
        _put(ft, t.row, 28, o.sector_proxy)
        _put(ft, t.row, 29, o.excess_vs_sector, PCT)
        filled += 1

    # 3. Return Ledger — every row; returns as formulas over the row's own levels.
    lg = xl.create_sheet("Return Ledger")
    _title(lg, "ARTICLE-COMPANY RETURN LEDGER",
           "One row per (publication, company) and horizon. Returns are formulas over the levels "
           "in the same row; status says why a row has none. Only LONG (Overweight) is scored.")
    keys = [f.name for f in dataclasses.fields(Obs)]
    formula_cols = {"stock_return": ("entry_adj_close", "exit_adj_close"),
                    "world_return": ("world_entry", "world_exit"),
                    "sp500_net_return": ("sp500_net_entry", "sp500_net_exit"),
                    "sector_return": ("sector_entry", "sector_exit")}
    excess_cols = {"excess_vs_world": "world_return", "excess_vs_sp500": "sp500_net_return",
                   "excess_vs_sector": "sector_return"}
    _hdr(lg, 4, keys)
    col = {k: get_column_letter(i) for i, k in enumerate(keys, 1)}
    for i, o in enumerate(ledger, 5):
        d = dataclasses.asdict(o)
        for j, k in enumerate(keys, 1):
            v = d[k]
            if o.status == COMPLETE and k in formula_cols and v is not None:
                a, b = formula_cols[k]
                v = f"={col[b]}{i}/{col[a]}{i}-1"
            elif o.status == COMPLETE and k in excess_cols and v is not None:
                v = f"={col['stock_return']}{i}-{col[excess_cols[k]]}{i}"
            _put(lg, i, j, v, _fmt_for(k))
    lg.freeze_panes = "C5"

    # 4. Analysis sheets.
    for name, (title, rows) in sheets.items():
        table_sheet(xl, name, title.split(" || ")[0],
                    title.split(" || ")[1] if " || " in title else "", rows)

    # 5. Run summary, placed right after START HERE.
    sm = xl.create_sheet("Audit Run Summary", 1)
    _title(sm, "AUTOMATED AUDIT RUN | POLYGON + CIQ", "Generated by vinci.research_audit; every "
           "figure below traces to 'Return Ledger', 'Polygon Daily Prices' and 'CIQ Daily Data'.")
    sm.column_dimensions["A"].width = 44
    sm.column_dimensions["B"].width = 120
    for i, (k, v) in enumerate(summary_lines, 4):
        _put(sm, i, 1, k, bold=True)
        _put(sm, i, 2, v)
        sm.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    xl.calculation.fullCalcOnLoad = True
    xl.save(dst)
    return {"filled": filled, "pilot_untouched": untouched_pilot, "not_computed": not_computed,
            "price_rows": r - 5}


def write_json(path: Path, obj: Any) -> None:
    path.write_text(json.dumps(obj, indent=2, default=_cell))


def identity_rows(ids: dict[str, Identity], series: dict[str, Series]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for t, i in sorted(ids.items()):
        s = series.get(i.provider_ticker) if i.provider_ticker else None
        out.append({
            "audit_ticker": t, "company_in_workbook": i.company, "status": i.status,
            "provider": "Polygon" if i.provider_ticker else "",
            "provider_ticker": i.provider_ticker,
            "security_identity": (f"{i.provider_name} | CIK {i.cik} | FIGI {i.composite_figi} / "
                                  f"{i.share_class_figi} | type {i.security_type}"
                                  if i.provider_ticker else ""),
            "currency": i.currency or i.listing_currency, "exchange": i.listing_mic,
            "list_date": i.list_date, "metadata_as_of": i.as_of, "sic": i.sic_code,
            "sic_description": i.sic_description, "sector": i.sector, "industry": i.industry,
            "first_bar": s.points[0].day if s and s.points else "",
            "last_bar": s.points[-1].day if s and s.points else "",
            "n_bars": len(s.points) if s else 0,
            "first_live_bar": s.first_live_day if s else "",
            "cache_live_overlap": (f"{s.overlap.n_overlap} sessions, {s.overlap.n_mismatch} "
                                   f"mismatches, max rel diff {s.overlap.max_rel_diff:.2e}"
                                   if s else ""),
            "dividends_applied": len(s.dividends_used) if s else 0,
            "splits_applied": "; ".join(f"{x.execution_date} {x.split_to:g}:{x.split_from:g}"
                                        for x in s.splits_used) if s else "",
            "session_gaps": len(s.gaps) if s else 0,
            "blocks": "; ".join(b.reason for b in s.blocks) if s else "",
            "warnings": " | ".join(s.warnings) if s else "",
            "detail": i.detail,
            "polygon_lines_found_not_substituted": " || ".join(i.alternatives),
        })
    return out
