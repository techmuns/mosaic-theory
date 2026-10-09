"""Post-recalculation QA of the updated workbook — the spreadsheet must say what the ledger says.

Run AFTER a formula engine (LibreOffice) has recalculated the copy: openpyxl writes formulas with
no cached values, so before recalculation every formula cell reads ``None`` and nothing here could
be checked. Three questions:
1. does each filled 'Forward Return Tests' row's OWN formula (R stock, S benchmark, T excess)
   compute the ledger's number from the linked cells?
2. are the CIQ pilot rows exactly as supplied (values compared against the input workbook)?
3. does any cell in any sheet hold a formula error?
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import openpyxl

from vinci.research_audit.workbook import SHEET_TESTS

_ERRORS = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
TOL = 1e-9


def _ledger(path: Path) -> dict[str, dict[str, str]]:
    with path.open() as fh:
        return {r["test_id"]: r for r in csv.DictReader(fh) if r["test_id"]}


def qa_workbook(xlsx: Path, input_xlsx: Path, ledger_csv: Path) -> list[tuple[str, bool, str]]:
    out: list[tuple[str, bool, str]] = []
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    src = openpyxl.load_workbook(input_xlsx, data_only=True)
    led = _ledger(ledger_csv)
    ft, ft0 = wb[SHEET_TESTS], src[SHEET_TESTS]

    checked, bad = 0, []
    for r in range(5, ft.max_row + 1):
        tid = ft.cell(row=r, column=1).value
        o = led.get(str(tid)) if tid else None
        if o is None or o["status"] != "COMPLETE":
            continue
        if ft0.cell(row=r, column=14).value not in (None, ""):
            continue  # pilot row: checked below, never recomputed here
        got = [ft.cell(row=r, column=c).value for c in (18, 19, 20)]
        want = [float(o["stock_return"]), float(o["world_return"]), float(o["excess_vs_world"])]
        checked += 1
        if any(not isinstance(g, int | float) or abs(float(g) - w) > TOL
               for g, w in zip(got, want, strict=True)):
            bad.append(f"{tid}: sheet {got} vs ledger {want}")
    out.append((f"filled rows: sheet formulas R/S/T == ledger (tol {TOL:g})",
                checked > 0 and not bad, f"{checked} rows checked; {len(bad)} mismatches {bad}"))

    pilot_bad: list[str] = []
    n_pilot = 0
    for r in range(5, ft0.max_row + 1):
        if ft0.cell(row=r, column=14).value in (None, ""):
            continue
        n_pilot += 1
        for c in range(1, 25):
            a, b = ft0.cell(row=r, column=c).value, ft.cell(row=r, column=c).value
            if a != b and not (isinstance(a, float) and isinstance(b, float)
                               and abs(a - b) < 1e-12):
                pilot_bad.append(f"{ft.cell(row=r, column=c).coordinate}: {a!r} -> {b!r}")
    out.append(("CIQ pilot rows A:X unchanged from the input", n_pilot > 0 and not pilot_bad,
                f"{n_pilot} pilot rows; {len(pilot_bad)} differences {pilot_bad[:10]}"
                + (f" (showing 10 of {len(pilot_bad)})" if len(pilot_bad) > 10 else "")))

    errs: list[str] = []
    n_cells = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                n_cells += 1
                if isinstance(cell.value, str) and cell.value.strip() in _ERRORS:
                    errs.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    out.append(("no formula error values in any sheet", not errs,
                f"{n_cells} cells scanned; {len(errs)} errors {errs[:20]}"
                + (f" (showing 20 of {len(errs)})" if len(errs) > 20 else "")))

    orig_sheets = set(src.sheetnames)
    out.append(("every input sheet is present in the updated copy",
                orig_sheets <= set(wb.sheetnames),
                f"missing: {sorted(orig_sheets - set(wb.sheetnames))}"))
    return out


def append_to_log(log: Path, results: list[tuple[str, bool, str]], recalc: dict[str, Any]) -> None:
    lines = ["", "## 11. Workbook QA (after LibreOffice recalculation)", "",
             f"Recalculation: status `{recalc.get('status')}`, formulas "
             f"{recalc.get('total_formulas')}, errors {recalc.get('total_errors')}.", "",
             "| check | result | detail |", "|---|---|---|"]
    lines += [f"| {c} | {'PASS' if ok else 'FAIL'} | {d} |" for c, ok, d in results]
    with log.open("a") as fh:
        fh.write("\n".join(lines) + "\n")
