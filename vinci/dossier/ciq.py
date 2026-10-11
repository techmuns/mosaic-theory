"""CIQ parsing — file-check → format-branch, per ``docs/ciq_extraction_notes.md``.

The cardinal rule from the notes: **CIQ extensions mislead; verify the real format before parsing.**
We sniff header bytes (portable + testable) rather than trust the extension, then branch:
``.rtf`` real-RTF → ``striprtf``; OLE-disguised ``.rtf``/``.doc`` transcript → needs LibreOffice
(fail loud, never the ~12-word ``striprtf`` garbage); MIME-wrapped ``.doc`` → quoted-printable +
BeautifulSoup; legacy ``.xls`` BIFF → ``xlrd``; ``.xlsx`` → ``openpyxl``; ``.pdf`` → ``pdftotext``.

``xlrd``/``openpyxl``/``pandas`` are imported lazily inside their branch and fail loud with an
actionable message if absent (they're optional until real CIQ data lands). Every extractor fails
loud to UNKNOWN — a missing or unrecognised cell is never a fabricated number.
"""

from __future__ import annotations

import enum
import html as _html
import quopri
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

_OLE_MAGIC = b"\xd0\xcf\x11\xe0"
_ZIP_MAGIC = b"PK\x03\x04"
_MIN_RTF_WORDS = 50  # below this from a multi-KB file ⇒ striprtf hit an OLE-disguised doc

#: CIQ's OWN WORDS FOR "NOTHING TO LIST" — the difference between a parser that failed and a view
#: that is genuinely empty. A relationships export for a company with no disclosed counterparties is
#: a large RTF carrying a couple of sentences, which trips the thin-file guard above; before this,
#: BURL's customers view raised CiqParseError and the page reported customer concentration UNKNOWN,
#: blaming "an OLE-disguised Word doc" it had never checked for. The file was real RTF, parsed
#: perfectly, and said: "There are no recently disclosed customers listed".
#:
#: SCOPE, MEASURED (2026-09-28, the 26 ticker folders under data/manual/ciq/): BURL is the only
#: name whose customers view is empty; the other 25 parse rows, ROST and TGT among them. So this
#: is NOT "every retailer" — a company sells to disclosed businesses or it does not, and BURL
#: happens not to (0 customer rows beside 33 supplier rows in the same bundle). What makes the
#: case worth handling is not its frequency but its silence: the export is well-formed, the
#: reader worked, and the only thing wrong was our account of why the answer was empty.
#:
#: The sentinel is CIQ's phrasing, matched loosely enough to survive a wording change and tightly
#: enough that an OLE-disguise (which yields punctuation and fragments, not sentences) cannot fake
#: it. A thin file WITHOUT it still raises, which is the case the guard was written for.
_EMPTY_VIEW_SENTINEL = re.compile(r"there (?:are|were) no\b[^.\n]{0,80}\b(?:listed|disclosed)",
                                  re.IGNORECASE)


def is_empty_view(text: str) -> bool:
    """True when a CIQ view says IN ITS OWN WORDS that it has nothing to list.

    THE ONE READER OF ``_EMPTY_VIEW_SENTINEL``. Two callers need this question answered — the
    thin-RTF guard below (is a short file broken, or just empty?) and ``ciq_sheets.customers``
    (did the table parser fail, or was there no table?) — and two regexes over one phrasing is the
    divergent-parser bug class the house rule names. They read the same sentence through here."""
    return _EMPTY_VIEW_SENTINEL.search(text) is not None


class CiqFormat(enum.StrEnum):
    """The REAL format behind a CIQ file (sniffed from bytes, not the extension)."""

    BIFF_XLS = "biff_xls"  # legacy Excel binary → xlrd
    OOXML = "ooxml"  # .xlsx/.docx zip → openpyxl
    RTF_TEXT = "rtf_text"  # real RTF → striprtf
    OLE_DISGUISED = "ole_disguised"  # .doc/.rtf that is really an OLE Word doc → needs LibreOffice
    MIME_DOC = "mime_doc"  # MIME-wrapped quoted-printable HTML filing → quopri + bs4
    HTML = "html"
    PDF = "pdf"
    UNKNOWN = "unknown"


class CiqParseError(RuntimeError):
    """A CIQ file could not be parsed in this environment (e.g. a needed library/tool is absent)."""


def classify(path: Path) -> CiqFormat:
    """Sniff the real format from header bytes (the notes' 'never trust the extension' rule)."""
    head = path.read_bytes()[:512]
    ext = path.suffix.lower()
    if head[:4] == _OLE_MAGIC:
        # OLE compound doc: a legacy .xls IS BIFF; an OLE-backed .rtf/.doc is a disguised Word doc.
        return CiqFormat.BIFF_XLS if ext == ".xls" else CiqFormat.OLE_DISGUISED
    if head[:4] == _ZIP_MAGIC:
        return CiqFormat.OOXML
    if head[:4] == b"%PDF":
        return CiqFormat.PDF
    stripped = head.lstrip()
    if stripped[:5] == b"{\\rtf":
        return CiqFormat.RTF_TEXT
    low = stripped[:200].lower()
    if low[:6] == b"<html>" or low[:9] == b"<!doctype" or low[:5] == b"<html":
        return CiqFormat.HTML
    if any(m in head[:200] for m in (b"MIME-Version", b"X-Sender", b"Content-Type", b"------=")):
        return CiqFormat.MIME_DOC
    return CiqFormat.UNKNOWN


def read_text(path: Path, fmt: CiqFormat | None = None) -> str:
    """Decode a text-bearing CIQ file to clean prose. Fail loud if a needed tool is absent."""
    fmt = fmt or classify(path)
    if fmt is CiqFormat.RTF_TEXT:
        from striprtf.striprtf import rtf_to_text

        raw_text = path.read_text(encoding="utf-8", errors="replace")
        text = rtf_to_text(raw_text)  # type: ignore[no-untyped-call]
        if (path.stat().st_size > 4096 and len(text.split()) < _MIN_RTF_WORDS
                and not is_empty_view(text)):
            raise CiqParseError(
                f"{path.name}: striprtf yielded {len(text.split())} words from a "
                f"{path.stat().st_size}B file — this is an OLE-disguised Word doc, not real RTF; "
                "pre-convert with LibreOffice/Word (see ciq_extraction_notes.md)"
            )
        return str(text)
    if fmt in (CiqFormat.MIME_DOC, CiqFormat.HTML):
        from bs4 import BeautifulSoup

        raw = path.read_bytes()
        if fmt is CiqFormat.MIME_DOC:
            # strip MIME headers (up to the blank line), then decode quoted-printable
            sep = b"\r\n\r\n" if b"\r\n\r\n" in raw else b"\n\n"
            raw = quopri.decodestring(raw.split(sep, 1)[-1])
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return _html.unescape(soup.get_text(" ", strip=True))
    if fmt is CiqFormat.PDF:
        import shutil
        import subprocess

        if not shutil.which("pdftotext"):
            raise CiqParseError(f"{path.name}: pdftotext not installed — cannot read PDF")
        out = subprocess.run(  # noqa: S603 — fixed args, path from the local manual store
            ["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, check=True
        )
        return out.stdout
    if fmt is CiqFormat.OLE_DISGUISED:
        import shutil
        import subprocess
        import tempfile

        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice is None:
            raise CiqParseError(
                f"{path.name}: OLE-disguised Word doc — needs LibreOffice; soffice not installed "
                "(ciq_extraction_notes.md: striprtf returns ~12 words of garbage)"
            )
        with tempfile.TemporaryDirectory() as td:
            subprocess.run(  # noqa: S603 — fixed args, path from the local manual store
                [soffice, "--headless", "--convert-to", "txt", "--outdir", td, str(path)],
                capture_output=True, check=True,
            )
            txt_path = Path(td) / f"{path.stem}.txt"
            if not txt_path.exists():
                raise CiqParseError(f"{path.name}: LibreOffice conversion produced no .txt output")
            return txt_path.read_text(encoding="utf-8", errors="replace")
    raise CiqParseError(f"{path.name}: format {fmt.value} is not text — use read_sheets()")


def read_sheets(path: Path, fmt: CiqFormat | None = None) -> dict[str, list[list[Any]]]:
    """Read a CIQ spreadsheet into ``{sheet: rows}``. Lazy xlrd/openpyxl; fail loud if absent."""
    fmt = fmt or classify(path)
    if fmt is CiqFormat.BIFF_XLS:
        try:
            import xlrd  # type: ignore[import-untyped]
        except ImportError as exc:
            raise CiqParseError(
                f"{path.name}: legacy .xls needs `xlrd` (pip install xlrd) — not installed"
            ) from exc
        book = xlrd.open_workbook(str(path))
        return {s.name: [s.row_values(r) for r in range(s.nrows)] for s in book.sheets()}
    if fmt is CiqFormat.OOXML:
        try:
            import openpyxl  # type: ignore[import-untyped]
        except ImportError as exc:
            raise CiqParseError(
                f"{path.name}: .xlsx needs `openpyxl` (pip install openpyxl) — not installed"
            ) from exc
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        return {
            ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets
        }
    raise CiqParseError(f"{path.name}: format {fmt.value} is not a spreadsheet")


# --- workbook-aware model: a CIQ export is a multi-sheet workbook; address = (kind, sheet) ----
#
# A CIQ Financials export is ONE 13-subsheet workbook (Income Statement / Balance Sheet / Cash
# Multiples / Segments / Capital Structure Details / Ratios / …); the Estimates export is its own
# workbook (Consensus / Recent Changes / Multiples-forward / Surprise / Revisions). A module is an
# explicit (workbook-kind, sheet) ADDRESS — never module→file by filename keyword, which silently
# grabs the WRONG Multiples (the historical Financials one vs the forward Estimates one differ).

# Filename → workbook KIND (the export type). Financials may appear twice (Annual + Quarterly),
# disambiguated by a verbatim Period-Type read, not the filename.
_WORKBOOK_KEYWORDS: dict[str, tuple[str, ...]] = {
    "estimates": ("estimate", "consensus"),
    "financials": ("financial", "income", "balance", "cash flow", "cashflow"),
    "customers": ("customer", "supplier", "business relationship"),
    "transcript": ("earnings call", "transcript", "conference call"),
    "key_developments": ("key development",),
    "profile": ("company profile", "tearsheet"),
    "comps": ("comparable", "comps"),  # "Company Comparable Analysis <name>.xls" (peer comps)
    # CIQ Charting export (2026-09-03). Sheets are `Pane 1` / `Raw` / `Attributions` — NOT the
    # Financials/Estimates layout — and the header row VARIES by chart type, so nothing here may
    # hardcode a position (notes lines 72, 361). Unregistered until now, which meant
    # `CiqBundle.from_files` DROPPED a charting export silently.
    "charting": ("charting", "chart export"),
    # S&P-MI Credit Health Panel, exported ON ITS OWN (2026-10-05). The same relative score as the
    # comps workbook's "Credit Health Panel" sheet, ranked within the name's GICS INDUSTRY (21–40
    # public companies) instead of the 10–11-name comp set. Unregistered until now, so ten folders'
    # exports were dropped without a word — the `charting` failure, a second time.
    "credit_health": ("credit health",),
}
# unparseable / withheld CIQ cells — ALL mean UNAVAILABLE (notes lines 61, 74); never coerce to 0.
_UNAVAILABLE_CELLS = frozenset({"", "-", "—", "NM", "NA", "N/A", "Entitlement Needed"})


class CiqUnavailableError(RuntimeError):
    """A needed CIQ workbook/subsheet is absent or a parse tool missing — fail loud, never fake."""


class CiqAmbiguousError(RuntimeError):
    """A name's folder holds MORE THAN ONE export of a kind the reader takes ONE of (2026-10-11).

    THE LOADER PICKED THE FIRST BY SORT ORDER, AND THE STALE FILE SORTS FIRST. CIQ renamed a
    same-kind export between pulls in 3 of the 6 archived re-exports — BRBR's and UBER's comps
    went "…Inc .xls" → "…Inc.xls", GOOGL's Key Developments `.xls` → `.rtf` — so a re-export copied
    in beside the old file leaves two, and the space sorts before the dot: the August comps would
    have priced UBER with nothing said. Only the archive-then-mirror sync prevented it, twice.

    NOT A `CiqUnavailableError`, ON PURPOSE. Every reader catches that one and degrades to a
    missing section; an ambiguous folder is not an absent export, and a page assembled on a guess
    between two exports is exactly what this exists to stop. It propagates and stops the name."""


#: Kinds a folder LEGITIMATELY holds several of — read whole through
#: :meth:`CiqBundle.all_workbooks`, and exempt from the one-export rule by name, never by a
#: heuristic. `customers` matches BOTH of CIQ's business-relationship exports (`… Customers.rtf`
#: and `… Suppliers.rtf`, on every name — its one reader picks the Customers file by name);
#: `transcript` holds one file per earnings call.
MULTI_FILE_KINDS = frozenset({"customers", "transcript"})

_AMBIGUITY_REMEDY = ("keep ONE export per kind: move the stale one(s) to "
                     "data/manual/ciq-archive/replaced-<YYYY-MM-DD>/<TICKER>/ and re-run")


def _customers_role(name: str) -> str | None:
    """Which of the two business-relationship exports a `customers`-kind file is, by name."""
    low = name.lower()
    return "customers" if "customer" in low else "suppliers" if "supplier" in low else None


def ambiguous_exports(workbooks: tuple[CiqWorkbook, ...] | list[CiqWorkbook]) -> list[str]:
    """Every single-file kind (by period type, for financials) the workbooks hold more than once,
    each with the files it found — empty when every kind is held once.

    `customers` is two exports, not one: the Customers file and the Suppliers file are each held
    ONCE. Its reader picks the Customers file BY NAME, first match, so a re-pulled Customers export
    beside the old one (`.xls` → `.rtf`, the GOOGL Key Developments rename) is the same stale-wins
    shape as any other kind. Transcripts alone are exempt — one per call."""
    groups: dict[tuple[str, str | None], list[str]] = {}
    for w in workbooks:
        if w.kind == "customers":
            groups.setdefault((w.kind, _customers_role(w.path.name)), []).append(w.path.name)
        elif w.kind not in MULTI_FILE_KINDS:
            groups.setdefault((w.kind, w.freq), []).append(w.path.name)
    return [f"{kind}{f' ({freq})' if freq else ''}: {sorted(names)}"
            for (kind, freq), names in sorted(groups.items(),
                                               key=lambda kv: (kv[0][0], str(kv[0][1])))
            if len(names) > 1]


def clean_num(cell: Any) -> float | None:
    """A CIQ numeric cell → float, or None when UNAVAILABLE (empty / '-' / 'NM' / 'Entitlement'
    / 'NA') — never coerced to 0 or a value (notes lines 61, 74)."""
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    s = str(cell).strip()
    if s in _UNAVAILABLE_CELLS:
        return None
    try:
        return float(s.replace(",", "").replace("%", ""))
    except ValueError:
        return None


def excel_date(cell: Any) -> date | None:
    """A CIQ date cell → date: Excel serial (Multiples/Estimates), already-decoded datetime, or an
    ISO/`Mon-DD-YYYY` string (notes lines 63-66). UNAVAILABLE → None."""
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=int(cell))).date()
        except (ValueError, OverflowError):
            return None
    s = str(cell).strip()
    for fmt in ("%Y-%m-%d", "%b-%d-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def period_type(income_rows: list[list[Any]]) -> str | None:
    """Verbatim Period-Type read: scan for the "Period Type:" label cell, return the next cell
    (notes lines 136-138 — ground frequency on this cell, NEVER on column-count/shape inference)."""
    for row in income_rows[:14]:
        for j, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip() == "Period Type:" and j + 1 < len(row):
                return str(row[j + 1]).strip().lower() or None
    return None


def _financials_freq(name: str, income_rows: list[list[Any]]) -> str | None:
    """The period type of a financials export. The FILENAME is authoritative — CIQ names them
    '…Financials_Quarterly' / '…Financials_Annual', so we match the '_quarterly' / '_annual' token
    (anchored on the underscore so a company NAME with 'quarterly'/'annual' can't false-match).
    :func:`period_type` reads the Income Statement's "Period Type:" cell, which is unreliable when
    that sheet is EMPTY — AORT's Quarterly export has an empty Income Statement whose boilerplate
    reads "Annual", so a content-only read mis-tagged the quarterly workbook 'annual' and a
    freq="quarterly" request silently fell back to the Annual workbook. Fall back to the verbatim
    cell only for an un-named export."""
    low = name.lower()
    if "_quarterly" in low:
        return "quarterly"
    if "_annual" in low:
        return "annual"
    return period_type(income_rows)


@dataclass(frozen=True)
class CiqWorkbook:
    kind: str
    path: Path
    fmt: CiqFormat
    freq: str | None  # "annual"/"quarterly" for financials (verbatim Period-Type), else None


@dataclass
class CiqBundle:
    """A name's CIQ exports as workbooks. A field resolves to a (workbook-kind, sheet) address via
    :meth:`sheet` (fail-loud UNAVAILABLE). Adding a consumer = register an address + write its
    extractor — no loader re-architecture; unconsumed sheets are reachable but not pre-extracted."""

    ticker: str
    workbooks: tuple[CiqWorkbook, ...]
    _cache: dict[Path, dict[str, list[list[Any]]]] = field(default_factory=dict, repr=False)
    #: THE COMPS PANEL'S REAL DATA DATE, injected by the assembler — never derived here.
    #: CIQ's `As-Of Date:` header is the EXPORT stamp; the data is an earlier trading close, so
    #: every consumer of `comps_as_of_date` would otherwise publish a date the data does not reach.
    #: Resolving it needs PRICE BARS, which this module has no business knowing about — so the
    #: assembler resolves it once (it holds both the bundle and the store) and sets it here, and
    #: `ciq_sheets` RENDERS what it is given. One slot corrects all eight comps consumers without
    #: threading price data through a workbook parser.
    #: PER-ASSEMBLY BY CONSTRUCTION: a bundle is built per ticker per run, so this cannot leak
    #: between names. Left None (the default) nothing changes and the header stamp stands.
    resolved_comps_as_of: date | None = field(default=None, repr=False)
    #: Files in the folder that match NO workbook kind, by name. Reported by the assembler rather
    #: than skipped: a dropped export is a fact the operator went to the trouble of producing, and
    #: twice now (`charting`, then ten Credit Health exports) one vanished with no trace.
    unrecognised: tuple[str, ...] = ()

    @classmethod
    def from_files(cls, ticker: str, paths: list[Path]) -> CiqBundle:
        wbs: list[CiqWorkbook] = []
        cache: dict[Path, dict[str, list[list[Any]]]] = {}
        unrecognised: list[str] = []
        for p in paths:
            # Sync and editor droppings (`.DS_Store`, `._<name>` AppleDouble files, `~$book.xls`
            # lock files) are not exports — skipped BEFORE classification, because a lock file
            # carries its workbook's name and would otherwise count as a second export of its kind
            # (an Excel window open on the comps would refuse the whole folder).
            if p.name.startswith((".", "~$")):
                continue
            name = p.name.lower()
            kind = next(
                (k for k, kws in _WORKBOOK_KEYWORDS.items() if any(w in name for w in kws)), None
            )
            if kind is None:
                unrecognised.append(p.name)
                continue
            try:
                fmt = classify(p)
            except OSError:
                fmt = CiqFormat.UNKNOWN
            # The FILENAME token names the period type even when the workbook cannot be read
            # (no `xlrd`, a corrupt or non-spreadsheet export): an unreadable Annual + Quarterly
            # pair is two period types, not two exports, and the parse failure surfaces — as it
            # always did — when a reader opens the sheet.
            freq = _financials_freq(name, []) if kind == "financials" else None
            if kind == "financials" and fmt in (CiqFormat.BIFF_XLS, CiqFormat.OOXML):
                try:
                    cache[p] = read_sheets(p, fmt)
                    freq = _financials_freq(name, cache[p].get("Income Statement", []))
                except CiqParseError:
                    pass
            wbs.append(CiqWorkbook(kind=kind, path=p, fmt=fmt, freq=freq))
        # REFUSED AT CONSTRUCTION, before any reader runs: several readers sit inside broad
        # handlers that turn a raise into a missing section, so the one place a guess between two
        # exports cannot become a quiet page is here.
        if dups := ambiguous_exports(wbs):
            raise CiqAmbiguousError(
                f"{ticker}: the CIQ folder holds more than one export of a kind read as one — "
                f"{'; '.join(dups)}. {_AMBIGUITY_REMEDY}")
        return cls(ticker=ticker, workbooks=tuple(wbs), _cache=cache,
                   unrecognised=tuple(unrecognised))

    def workbook(self, kind: str, *, freq: str | None = None) -> CiqWorkbook | None:
        """The ONE export of ``kind`` (of period type ``freq``, for financials), or None.

        NEVER A PICK BETWEEN SEVERAL (2026-10-11). This returned the first candidate by sort
        order, and with ``freq`` it fell back to the first of ANY period type — the shape that once
        served AORT's annual workbook to a quarterly request. Now several candidates refuse
        (:class:`CiqAmbiguousError`). A LONE export still answers, as before — WHATEVER its period
        type: a folder holding only an Annual export answers a quarterly request with it (kept
        deliberately; no live folder is in that state). `MULTI_FILE_KINDS` keep their old answer
        — they are read through :meth:`all_workbooks`."""
        cands = [w for w in self.workbooks if w.kind == kind]
        if not cands:
            return None
        if kind in MULTI_FILE_KINDS:
            if freq is not None:
                return next((w for w in cands if w.freq == freq), None) or cands[0]
            return cands[0]
        names = sorted(w.path.name for w in cands)
        if freq is not None:
            match = [w for w in cands if w.freq == freq]
            if len(match) == 1:
                return match[0]
            if len(cands) == 1:
                return cands[0]
            raise CiqAmbiguousError(
                f"{self.ticker}: {kind} ({freq}) — "
                + (f"{len(match)} exports of that period type: {sorted(w.path.name for w in match)}"
                   if match else f"none of that period type among {names}, and another period "
                   "type is never substituted among several")
                + f". {_AMBIGUITY_REMEDY}")
        if len(cands) > 1:
            raise CiqAmbiguousError(
                f"{self.ticker}: {kind} — {len(cands)} exports {names}, read as one. "
                f"{_AMBIGUITY_REMEDY}")
        return cands[0]

    def all_workbooks(self, kind: str) -> tuple[CiqWorkbook, ...]:
        """Every workbook of ``kind`` (vs :meth:`workbook`, which returns one) — for the
        earnings-call transcripts, where a name commonly has more than one quarter on file."""
        return tuple(w for w in self.workbooks if w.kind == kind)

    def sheet(self, kind: str, sheet_name: str, *, freq: str | None = None) -> list[list[Any]]:
        """Parsed rows at a (kind, sheet) address, or fail loud :class:`CiqUnavailableError`."""
        wb = self.workbook(kind, freq=freq)
        if wb is None:
            raise CiqUnavailableError(f"no '{kind}' CIQ workbook for {self.ticker}")
        if wb.path not in self._cache:
            try:
                self._cache[wb.path] = read_sheets(wb.path, wb.fmt)
            except CiqParseError as exc:
                raise CiqUnavailableError(str(exc)) from exc
        sheets = self._cache[wb.path]
        if sheet_name not in sheets:
            raise CiqUnavailableError(f"{kind} workbook {wb.path.name} has no '{sheet_name}' sheet")
        return sheets[sheet_name]

    def has(self, kind: str) -> bool:
        return any(w.kind == kind for w in self.workbooks)

    @property
    def is_empty(self) -> bool:
        return not self.workbooks
