"""vinci.research_audit — the sell-side historical alpha audit.

Fixtures are REAL: bars and the session calendar come from Polygon responses captured 2026-10-09
(``tests/fixtures/research_audit/polygon_real_bars.json`` names its sources); workbook strings
are verbatim rows of the operator's audit workbook. Each guard ships a positive AND a negative pin.
"""

from __future__ import annotations

import dataclasses
import json
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from vinci.research_audit import aggregate as ag
from vinci.research_audit.benchmarks import Benchmark, currency_fit
from vinci.research_audit.classify import (
    Availability,
    Direction,
    UnclassifiedError,
    direction_of,
    listing_of,
    role_of,
    signal_id_of,
)
from vinci.research_audit.engine import (
    COMPLETE,
    DUPLICATE_MENTION,
    INCOMPLETE_HORIZON,
    NO_PRICE_AT_EXIT,
    NOT_SCORED_NONDIRECTIONAL,
    UNAVAILABLE_POLYGON,
    EngineInputs,
    build_ledger,
    edate,
    session_after,
    session_on_or_after,
    verify_identity,
)
from vinci.research_audit.polygon import DailyBar, Dividend, GroupedDay, Split
from vinci.research_audit.prices import build_series, unadjust_cached
from vinci.research_audit.sec_sources import _sector_label, names_match
from vinci.research_audit.workbook import (
    MENTION_HEADERS,
    AuditWorkbook,
    Mention,
    Publication,
    WorkbookLayoutError,
    _table,
)

FIX = json.loads((Path(__file__).parent / "fixtures" / "research_audit"
                  / "polygon_real_bars.json").read_text())
SESSIONS = [date.fromisoformat(d) for d in FIX["sessions"]]
NVDA_SPLIT = Split(date(2024, 6, 10), split_from=1.0, split_to=10.0, polygon_id="nvda-2024")


def _bar(r: dict[str, Any], src: str = "polygon-grouped-cache") -> DailyBar:
    return DailyBar(date.fromisoformat(r["day"]), r["o"], r["h"], r["l"], r["c"], r["v"], src, "")


def _grouped(sym: str) -> dict[date, GroupedDay | None]:
    return {date.fromisoformat(r["day"]): GroupedDay(
        day=date.fromisoformat(r["day"]), fetched_at=date.fromisoformat(r["fetched"]),
        n_rows=12000, bars={sym: _bar(r)}) for r in FIX[sym]}


# --- timing -----------------------------------------------------------------------------------

def test_edate_clamps_to_month_end_like_excel() -> None:
    assert edate(date(2024, 1, 3), 6) == date(2024, 7, 3)  # the CIQ pilot's own target
    assert edate(date(2024, 1, 31), 1) == date(2024, 2, 29)
    assert edate(date(2024, 8, 31), 6) == date(2025, 2, 28)
    assert edate(date(2024, 6, 5), 12) == date(2025, 6, 5)


def test_entry_is_the_session_after_publication_and_exit_rolls_past_closures() -> None:
    # Published Fri 2025-01-03 → entry Mon 2025-01-06 (never the publication day's own close).
    assert session_after(SESSIONS, date(2025, 1, 3)) == date(2025, 1, 6)
    # 2024-07-04 holiday and the 2025-01-09 market closure both roll to the next real session.
    assert session_on_or_after(SESSIONS, date(2024, 7, 4)) == date(2024, 7, 5)
    assert session_on_or_after(SESSIONS, date(2025, 1, 9)) == date(2025, 1, 10)
    # Negative pin: a session day maps to itself.
    assert session_on_or_after(SESSIONS, date(2025, 1, 8)) == date(2025, 1, 8)


# --- prices: splits, dividends, splice --------------------------------------------------------

def test_cached_bar_is_unadjusted_for_a_split_after_it_but_not_before() -> None:
    pre, on = (_bar(r) for r in FIX["NVDA"][1:3])  # 2024-06-07 (pre-split), 2024-06-10 (split day)
    fetched = date(2026, 8, 6)
    assert unadjust_cached(pre, fetched, [NVDA_SPLIT]).close == pytest.approx(pre.close * 10)
    assert unadjust_cached(on, fetched, [NVDA_SPLIT]).close == on.close
    # A split AFTER the response was fetched had not been applied to it → left alone.
    assert unadjust_cached(pre, date(2024, 6, 9), [NVDA_SPLIT]).close == pre.close


def test_real_split_neither_breaks_the_tr_chain_nor_passes_unrecorded() -> None:
    ok = build_series("NVDA", live=[], grouped=_grouped("NVDA"), dividends=[],
                      splits=[NVDA_SPLIT], sessions=None)
    p = ok.by_day
    assert p[date(2024, 6, 7)].close == pytest.approx(1208.88)  # the raw, as-traded close
    assert p[date(2024, 6, 10)].split_factor == 10.0
    tr = p[date(2024, 6, 10)].adj_close / p[date(2024, 6, 7)].adj_close - 1
    assert abs(tr) < 0.02 and not ok.blocks
    # Negative pin: the same raw prices WITHOUT the split record must block, never pass as a -90%.
    raw = {d: GroupedDay(d, date(2024, 6, 9), 1, {"NVDA": unadjust_cached(
        g.bars["NVDA"], date(2026, 8, 6), [NVDA_SPLIT])}) for d, g in _grouped("NVDA").items()
        if g is not None}
    bad = build_series("NVDA", live=[], grouped=dict(raw), dividends=[], splits=[], sessions=None)
    assert [b.reason.split()[0] for b in bad.blocks] == ["SUSPECTED_UNRECORDED_SPLIT"]
    assert bad.blocked_between(date(2024, 6, 7), date(2024, 6, 11))
    assert bad.blocked_between(date(2024, 6, 10), date(2024, 6, 11)) is None


def test_a_real_earnings_crash_is_a_warning_not_an_unrecorded_split() -> None:
    # TTD 2025-02-13: -33.0% after earnings, no split on record. It must not block (it did, as a
    # "0.670x split", in the first full run); the NVDA case above must still block.
    bars = [_bar(r, "live") for r in FIX["TTD_earnings_2025_02"]]
    s = build_series("TTD", live=bars, grouped={}, dividends=[], splits=[], sessions=None)
    assert not s.blocks
    assert any("2025-02-13" in w and "-33.0%" in w for w in s.warnings)


def test_dividend_is_reinvested_on_its_ex_date() -> None:
    div = Dividend(date(2024, 1, 4), 0.60, "USD", "CD", "2024-02-09", "axp")
    s = build_series("AXP", live=[], grouped=_grouped("AXP"), dividends=[div], splits=[],
                     sessions=None)
    a, b = s.by_day[date(2024, 1, 3)], s.by_day[date(2024, 1, 4)]
    assert b.div_cash == 0.60 and a.div_cash == 0.0
    assert b.adj_close / a.adj_close == pytest.approx((187.14 + 0.60) / 186.32)
    assert s.points[-1].adj_close == s.points[-1].close  # scaled to the last raw close
    # Negative pin: a non-USD dividend is never folded into a USD return.
    eur = dataclasses.replace(div, currency="EUR")
    s2 = build_series("AXP", live=[], grouped=_grouped("AXP"), dividends=[eur], splits=[],
                      sessions=None)
    assert s2.blocked_between(date(2024, 1, 3), date(2024, 1, 5)) == "NON_USD_DIVIDEND EUR"


def test_cache_live_overlap_must_reconcile() -> None:
    ov = FIX["AXP_overlap_2024_10_09"]
    day = date(2024, 10, 9)
    cache = {day: GroupedDay(day, date(2026, 8, 6), 12000, {"AXP": DailyBar(
        day, 1, 1, 1, ov["cache_close"], 1, "cache", "")})}
    live = [DailyBar(day, 1, 1, 1, ov["live_close"], 1, "live", "")]
    good = build_series("AXP", live=live, grouped=dict(cache), dividends=[], splits=[],
                        sessions=None)
    assert good.overlap.n_overlap == 1 and good.overlap.ok and not good.blocks
    off = [dataclasses.replace(live[0], close=ov["live_close"] * 1.01)]
    bad = build_series("AXP", live=off, grouped=dict(cache), dividends=[], splits=[],
                       sessions=None)
    assert bad.overlap.n_mismatch == 1 and bad.blocks and not bad.blocks[0].step


# --- classification ---------------------------------------------------------------------------

def test_only_overweight_is_long_and_unknown_text_raises() -> None:
    assert direction_of("Overweight") is Direction.LONG
    assert direction_of("Equal-Weight") is Direction.NEUTRAL
    assert direction_of("Underweight") is Direction.UNDERWEIGHT
    assert direction_of("Not stated") is Direction.NONDIRECTIONAL
    with pytest.raises(UnclassifiedError):
        direction_of("Buy")  # never mapped by resemblance
    with pytest.raises(UnclassifiedError):
        role_of("Top pick")
    assert role_of("Fabless industry example").value == "CONTEXTUAL"
    assert signal_id_of("MS-01", "Exhibit 55, report page 34", "Overweight") == "MS24_ADOPTER_OW"
    assert signal_id_of("MS-03", "Exhibit 79, report page 40", "Equal-Weight") == ""


def test_non_us_primary_listing_is_unavailable_and_never_substituted() -> None:
    assert listing_of("EXPN.L").availability is Availability.UNAVAILABLE_POLYGON
    assert listing_of("601138.SS").exchange_mic == "XSHG"
    assert listing_of("AXP").availability is Availability.POLYGON_CANDIDATE
    alt = [{"ticker": "EXPGY", "name": "EXPERIAN PLC S/ADR", "market": "otc", "type": "ADRC",
            "primary_exchange": "OTC Link"}]
    ident = verify_identity("EXPN.L", "Experian", None, None, None, alt)
    assert ident.status == UNAVAILABLE_POLYGON and ident.provider_ticker == ""
    assert "EXPGY" in ident.alternatives[0]


def test_identity_name_check_fails_closed_unless_sec_confirms_the_cik() -> None:
    # Polygon's own record for MAR, verbatim — including its misspelling.
    det = {"ticker": "MAR", "name": "Marriot International Class A Common Stock",
           "primary_exchange": "XNAS", "type": "CS", "currency_name": "usd",
           "cik": "0001048286", "list_date": "1998-03-27", "sic_code": "7011"}
    plain = verify_identity("MAR", "Marriott International", det, date(2024, 6, 5),
                            date(2024, 6, 4))
    assert plain.status != "VERIFIED_POLYGON"
    good = verify_identity("MAR", "Marriott International", det, date(2024, 6, 5),
                           date(2024, 6, 4), sec=("0001048286", "MARRIOTT INTERNATIONAL INC /MD/"))
    assert good.status == "VERIFIED_POLYGON" and "SEC company_tickers" in good.detail
    wrong = verify_identity("MAR", "Marriott International", det, date(2024, 6, 5),
                            date(2024, 6, 4), sec=("0000000001", "MARRIOTT INTERNATIONAL INC"))
    assert wrong.status != "VERIFIED_POLYGON"


def test_nport_name_matching_is_strict() -> None:
    assert names_match("Williams Cos Inc/The", "Williams Companies")
    assert names_match("MSCI Inc", "MSCI")
    assert names_match("Marriott International Inc/MD", "Marriott International")
    assert not names_match("United Rentals Inc", "United Airlines Holdings")
    assert not names_match("Morgan Stanley", "Kinder Morgan")
    assert _sector_label("VANGUARD INFORMATION TECHNOLOGY INDEX FUND") == "Information Technology"


# --- engine end-to-end on real rows -----------------------------------------------------------

def _pub(pid: str, day: str) -> Publication:
    return Publication(pid, "Morgan Stanley", "Thematic equity / AI adopter survey", day, "Day",
                       day, "Mapping AI's Diffusion", "Edward Stanley; Todd Castagno",
                       "Edward Stanley / Todd Castagno", "Global", "claim", "testable", "16",
                       "names", "scope", "pdf", "caveat", "https://example.invalid/ms", "R!A5")


def _mention(pid: str, company: str, ticker: str, role: str, rating: str, place: str,
             row: int) -> Mention:
    return Mention(pid, "Morgan Stanley", company, ticker, "US", role, rating, "Analyst", place,
                   "", "", f"Company Mentions!A{row}")


def _inputs(mentions: tuple[Mention, ...], last: date, drop: date | None = None) -> EngineInputs:
    wb = AuditWorkbook(Path("x.xlsx"), "sha", (), (_pub("MS-01", "2024-07-02"),), (), mentions,
                       (), (), "", "", {})
    # The calendar runs PAST `last`: only the last-completed-session guard may refuse a row.
    sess = [d for d in SESSIONS if d >= date(2024, 7, 1)]
    bars = [DailyBar(d, 100, 100, 100, 100 + i, 1000, "live", "") for i, d in enumerate(sess)
            if d != drop]
    s = build_series("AXP", live=bars, grouped={}, dividends=[], splits=[], sessions=sess)
    lv = {d: 100.0 + i / 2 for i, d in enumerate(sess)}
    bench = Benchmark("URTH_TR_PROXY", "URTH", "PROXY", lv, "test")
    det = {"ticker": "AXP", "name": "American Express Company", "primary_exchange": "XNYS",
           "type": "CS", "currency_name": "usd", "cik": "0000004962", "list_date": "1972-06-01"}
    ids = {"AXP": verify_identity("AXP", "American Express", det, date(2024, 7, 3),
                                  date(2024, 7, 2)),
           "EXPN.L": verify_identity("EXPN.L", "Experian", None, None, None)}
    return EngineInputs(wb=wb, identities=ids, series={"AXP": s}, sessions=sess,
                        last_session=last, world=bench, ciq_world=bench, sp500=bench,
                        sector_proxies={})


AXP_OW = _mention("MS-01", "American Express", "AXP",
                  "AI adopter + existing Overweight stock rating", "Overweight",
                  "Exhibit 55, report page 34", 5)


def test_complete_row_returns_trace_to_bars_and_benchmark() -> None:
    led = build_ledger(_inputs((AXP_OW,), last=date(2025, 2, 28)))
    by = {o.months: o for o in led}
    o = by[6]
    # Published Tue 2024-07-02 → entry 2024-07-03; EDATE +6m = 2025-01-03, a real session.
    assert (o.status, o.entry_date, o.target_date, o.exit_date) == (
        COMPLETE, "2024-07-03", "2025-01-03", "2025-01-03")
    assert o.stock_return == pytest.approx(o.exit_adj_close / o.entry_adj_close - 1)
    assert o.excess_vs_world == pytest.approx(o.stock_return - o.world_return)
    assert o.world_benchmark == "URTH_TR_PROXY"
    # 9m / 12m targets lie beyond the last completed session → not computed.
    assert by[9].status == by[12].status == INCOMPLETE_HORIZON
    assert by[9].stock_return is None and by[12].stock_return is None


def test_no_carry_forward_across_a_missing_exit_bar() -> None:
    led = build_ledger(_inputs((AXP_OW,), last=date(2025, 2, 28), drop=date(2025, 1, 3)))
    o = next(x for x in led if x.months == 6)
    assert o.status == NO_PRICE_AT_EXIT and o.stock_return is None
    # Negative pin: with the bar present the same row completes.
    assert next(x for x in build_ledger(_inputs((AXP_OW,), last=date(2025, 2, 28)))
                if x.months == 6).status == COMPLETE


def test_incomplete_horizon_is_never_computed() -> None:
    led = build_ledger(_inputs((AXP_OW,), last=date(2024, 12, 31)))
    assert {o.status for o in led} == {INCOMPLETE_HORIZON}
    assert all(o.stock_return is None for o in led)


def test_mentions_are_never_converted_to_recommendations() -> None:
    now_ctx = _mention("MS-01", "American Express", "AXP",
                       "Named in AI-agent infrastructure/vendor mapping (not itself an explicit "
                       "buy)", "Not stated", "Exhibit 6, report page 2", 6)
    exp = _mention("MS-01", "Experian", "EXPN.L", "AI adopter + existing Overweight stock rating",
                   "Overweight", "Exhibit 55, report page 34", 7)
    led = build_ledger(_inputs((now_ctx, AXP_OW, exp), last=date(2024, 7, 12)))
    by = {(o.mention_src, o.months): o for o in led}
    # The unrated row of a company that is ALSO rated in the same publication is the duplicate;
    # the rated row is the one counted.
    assert by[("Company Mentions!A6", None)].status == DUPLICATE_MENTION
    assert by[("Company Mentions!A5", 6)].direction == "LONG"
    assert {o.status for o in led if o.audit_ticker == "EXPN.L"} == {UNAVAILABLE_POLYGON}
    lone = build_ledger(_inputs((now_ctx,), last=date(2024, 7, 12)))
    assert [o.status for o in lone] == [NOT_SCORED_NONDIRECTIONAL]
    assert lone[0].stock_return is None


# --- aggregation ------------------------------------------------------------------------------

def _obs(direction: str, excess: float, pub: str = "MS-01", tic: str = "AXP") -> Any:
    o = build_ledger(_inputs((AXP_OW,), last=date(2024, 7, 12)))[0]
    return dataclasses.replace(o, status=COMPLETE, direction=direction, months=6,
                               excess_vs_world=excess, entry_date="2024-01-03",
                               exit_date="2024-07-03", pub_id=pub, audit_ticker=tic, flags=[])


def test_long_stats_never_pool_neutral_or_underweight_and_ties_are_not_wins() -> None:
    led = [_obs("LONG", 0.10, tic="A"), _obs("LONG", 0.0, tic="B"), _obs("NEUTRAL", -0.5, tic="C"),
           _obs("UNDERWEIGHT", 0.9, tic="D")]
    rows = {r["direction"]: r for r in ag.group_table(led, "publisher", lambda o: [o.publisher])}
    assert rows["LONG"]["6m_completed"] == 2
    assert rows["LONG"]["6m_win_rate_vs_world"] == 0.5
    assert rows["NEUTRAL"]["6m_completed"] == 1 and rows["UNDERWEIGHT"]["6m_completed"] == 1
    assert "SMALL_SAMPLE" in rows["LONG"]["flags"] and "SINGLE_ENTRY_DATE" in rows["LONG"]["flags"]
    board = ag.leaderboard(list(rows.values()))
    assert str(board[0]["rank_12m_median_excess"]).startswith("UNRANKED")


def test_currency_fit_finds_the_denomination_from_real_fx() -> None:
    fx = {date.fromisoformat(d): v for d, v in FIX["EURUSD"]}
    usd = {date.fromisoformat(d): v for d, v in FIX["URTH_raw_close"] if date.fromisoformat(d)
           in fx}
    proxy = Benchmark("URTH", "URTH", "PROXY", usd, "t")
    eur_series = Benchmark("CIQ", "CIQ", "INDEX", {d: v / fx[d] for d, v in usd.items()}, "t")
    fits = currency_fit(eur_series, proxy, {"C:EURUSD": fx})
    assert fits[0].hypothesis == "C:EURUSD" and fits[0].rms_21_session_gap < 1e-9
    usd_fit = currency_fit(proxy, proxy, {"C:EURUSD": fx})
    assert usd_fit[0].hypothesis.startswith("USD")


def test_workbook_header_drift_raises() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, h in enumerate(MENTION_HEADERS, 1):
        ws.cell(row=4, column=i, value=h)
    ws.cell(row=5, column=1, value="MS-01")
    assert _table(ws, 4, MENTION_HEADERS)[0][0] == 5
    ws.cell(row=4, column=7, value="Rating")  # a renamed column must not silently shift fields
    with pytest.raises(WorkbookLayoutError):
        _table(ws, 4, MENTION_HEADERS)


def test_mixed_tables_keep_every_column(tmp_path: Path) -> None:
    from vinci.research_audit.report import write_csv

    rows = [{"ticker": "AXP", "end_ciq": 2.16}, {"ticker": "AXP", "ex_date": "2024-01-04",
                                                  "cash": 0.6}]
    write_csv(tmp_path / "t.csv", rows)
    header = (tmp_path / "t.csv").read_text().splitlines()[0]
    assert header == "ticker,end_ciq,ex_date,cash"  # the second row's keys are not dropped


def test_rank_needs_two_cohorts_and_name_variants_are_flagged_not_merged() -> None:
    # Both spellings appear verbatim in the workbook's 'Named authors' column (MS-02 vs MS-01/03).
    base = {"level": "named author", "direction": "LONG", "flags": ""}
    rows = [{**base, "group": "Michelle Weaver", "12m_completed": 17, "12m_n_cohorts": 1,
             "12m_median_excess_vs_world": 0.10},
            {**base, "group": "Michelle M. Weaver", "12m_completed": 14, "12m_n_cohorts": 2,
             "12m_median_excess_vs_world": -0.07}]
    board = {r["group"]: r for r in ag.leaderboard(rows)}
    assert str(board["Michelle Weaver"]["rank_12m_median_excess"]).startswith("UNRANKED")
    assert board["Michelle M. Weaver"]["rank_12m_median_excess"] == 1
    assert "NAME_VARIANT_OF ['Michelle M. Weaver']" in board["Michelle Weaver"]["flags"]
    assert len(board) == 2  # flagged, never merged
